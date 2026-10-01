#!/usr/bin/env python3
"""Optional original template-based static Office export; never claims scene parity."""
import argparse,hashlib,json,pathlib,sys
import runtime
def main():
 p=argparse.ArgumentParser();p.add_argument('--input',required=True);p.add_argument('--template',required=True);p.add_argument('--out',required=True);p.add_argument('--kind',choices=runtime.KINDS,required=True);p.add_argument('--edition',choices=['metodologia','white-label'],default='metodologia');p.add_argument('--language',default='es');a=p.parse_args();src=pathlib.Path(a.input);template=pathlib.Path(a.template);out=pathlib.Path(a.out)
 if not template.is_file() or template.suffix.lower() not in ('.pptx','.docx'):p.error('Plantilla existente .pptx o .docx obligatoria')
 if out.suffix.lower()!=template.suffix.lower():p.error('Extensión debe coincidir con plantilla')
 if out.exists() or out.with_suffix(out.suffix+'.receipt.json').exists():p.error('Destino existente')
 if any(x.is_symlink() for x in (out,*out.parents)):p.error('Symlink rechazado')
 d=json.loads(src.read_text());errors=runtime.validate(d,a.kind,a.edition,src.parent)
 if errors:print(json.dumps({'status':'BLOCKED','errors':errors}));sys.exit(2)
 lang=a.language
 try:
  if template.suffix.lower()=='.pptx':
   from pptx import Presentation
   prs=Presentation(str(template));layouts=[layout for layout in prs.slide_layouts if len(layout.placeholders)>=2 and any(ph.is_placeholder and ph.placeholder_format.idx==0 for ph in layout.placeholders)]
   if not layouts:p.error('Plantilla no tiene layout con título y contenido')
   for s in d['sections']:
    slide=prs.slides.add_slide(layouts[0]);slide.shapes.title.text=runtime.localized(s['title'],lang);body=next((ph for ph in slide.placeholders if ph!=slide.shapes.title and ph.has_text_frame),None)
    if body is None:p.error('Layout no tiene placeholder de texto de contenido')
    body.text=runtime.localized(s.get('body'),lang)
    if s.get('prompt'):body.text+='\n\n'+runtime.default_prompt(s,lang)
   out.parent.mkdir(parents=True,exist_ok=True);prs.save(str(out))
  else:
   from docx import Document
   doc=Document(str(template))
   if 'Heading 1' not in doc.styles or 'Normal' not in doc.styles:p.error('Plantilla requiere estilos Heading 1 y Normal')
   doc.add_heading(runtime.localized(d['title'],lang),0)
   for s in d['sections']:
    doc.add_heading(runtime.localized(s['title'],lang),1);doc.add_paragraph(runtime.localized(s.get('body'),lang))
    if s.get('prompt'):doc.add_paragraph(runtime.default_prompt(s,lang))
   out.parent.mkdir(parents=True,exist_ok=True);doc.save(str(out))
 except ImportError as e:print('coverage_gap: dependencia opcional ausente: '+str(e),file=sys.stderr);sys.exit(3)
 receipt={'state':'RENDERED_DRAFT','inputSha256':hashlib.sha256(src.read_bytes()).hexdigest(),'templateSha256':hashlib.sha256(template.read_bytes()).hexdigest(),'outputSha256':hashlib.sha256(out.read_bytes()).hexdigest(),'limitations':['Static text export; no interactive SVG, timer, tabs, fields or animation parity.','Existing template content, themes and masters retained; append-only content.','Human visual review required.']};out.with_suffix(out.suffix+'.receipt.json').write_text(json.dumps(receipt,indent=2));print(json.dumps(receipt))
if __name__=='__main__':main()
